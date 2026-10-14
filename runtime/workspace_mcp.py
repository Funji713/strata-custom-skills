"""Opt-in Linux stdio MCP tools, confined to one bind-mounted workspace.

Paths are traversed using directory descriptors and O_NOFOLLOW, including every
parent directory. Existing files are replaced atomically; symlinks, hardlinks,
devices and paths outside the root are refused. No shell, deletion or network tools.
"""
from __future__ import annotations

import argparse
import hashlib
import io
import json
import os
import stat
import sys
import uuid
from contextlib import contextmanager

MAX_FILE = 256 * 1024
MAX_MESSAGE = 2 * 1024 * 1024
MAX_PPTX = 8 * 1024 * 1024


def schema(properties, required=()):
    return {"type": "object", "properties": properties, "required": list(required), "additionalProperties": False}


PATH = {"type": "string", "description": "Relative path inside /workspace, using /, e.g. src/main.py. Never a Windows or absolute path."}
OVERWRITE = {"type": "boolean", "default": False, "description": "Only set true when the user requested replacing an existing file."}
TOOLS = [
    {"name": "list_files", "description": "List one directory inside /workspace. Use this to find code or generated PPTX files.",
     "inputSchema": schema({"path": {"type": "string", "default": "."}})},
    {"name": "read_file", "description": "Read UTF-8 code/text from /workspace before editing. Returns numbered lines and SHA256. Read subsequent pages when truncated.",
     "inputSchema": schema({"path": PATH, "start_line": {"type": "integer", "minimum": 1, "default": 1},
                            "max_lines": {"type": "integer", "minimum": 1, "maximum": 200, "default": 100}}, ["path"])},
    {"name": "write_file", "description": "Actually save UTF-8 code/text into /workspace, creating parent directories. Existing files require overwrite=true. Returns saved path and SHA256.",
     "inputSchema": schema({"path": PATH, "content": {"type": "string", "maxLength": MAX_FILE}, "overwrite": OVERWRITE}, ["path", "content"])},
    {"name": "edit_file", "description": "Actually edit an existing UTF-8 code file: replace exactly one occurrence of old_text with new_text. Read first; optional expected_sha256 guards against stale edits.",
     "inputSchema": schema({"path": PATH, "old_text": {"type": "string", "minLength": 1}, "new_text": {"type": "string"},
                            "expected_sha256": {"type": "string", "pattern": "^[0-9a-f]{64}$"}}, ["path", "old_text", "new_text"])},
    {"name": "create_presentation", "description": "Actually create a real editable PowerPoint .pptx in /workspace. Use for PPT requests. Supply 1-40 slides with title and up to 6 short bullets each (max 360 characters total per slide). 16:9, Traditional Chinese font supported. Split dense content into more slides. Returns saved path; open it from the Windows Strata-workspace folder.",
     "inputSchema": schema({"path": {**PATH, "description": "Relative output path ending in .pptx, e.g. presentations/demo.pptx."},
                            "title": {"type": "string", "minLength": 1, "maxLength": 80}, "overwrite": OVERWRITE,
                            "slides": {"type": "array", "minItems": 1, "maxItems": 40, "items": schema({
                                "title": {"type": "string", "minLength": 1, "maxLength": 40, "description": "Short title, one paragraph without line breaks."},
                                "bullets": {"type": "array", "maxItems": 6, "items": {"type": "string", "maxLength": 120, "description": "One short paragraph without line breaks."}}}, ["title", "bullets"])}},
                           ["path", "title", "slides"])},
]


def text_arg(args, name, maximum=None, nonempty=False):
    value = args.get(name)
    if not isinstance(value, str) or (nonempty and not value) or (maximum and len(value) > maximum):
        raise ValueError(f"{name} must be {'nonempty ' if nonempty else ''}text" + (f" of at most {maximum} characters" if maximum else ""))
    # XML 1.0 (PPTX) and UTF-8 cannot encode invalid control characters/surrogates.
    if any((ord(c) < 32 and c not in "\n\r\t") or 0xD800 <= ord(c) <= 0xDFFF for c in value):
        raise ValueError(f"{name} contains an unsupported control character")
    return value


class Workspace:
    def __init__(self, root):
        if os.name != "posix" or not hasattr(os, "O_NOFOLLOW"):
            raise RuntimeError("Run this tool in the Linux Docker container; safe directory-descriptor access is required")
        self.root = os.open(root, os.O_RDONLY | os.O_DIRECTORY | os.O_NOFOLLOW)

    def close(self):
        os.close(self.root)

    @staticmethod
    def parts(path, directory=False):
        if directory and path == ".":
            return []
        if not isinstance(path, str) or not path or len(path) > 512 or any(c in path for c in ("\\", ":", "\0")):
            raise ValueError("Use a relative workspace path with / separators")
        parts = path.split("/")
        if any(p in ("", ".", "..") or len(p.encode("utf-8")) > 200 for p in parts):
            raise ValueError("Absolute paths, empty segments, . and .. are forbidden")
        if any(p.startswith(".strata-tmp-") for p in parts):
            raise ValueError("Internal temporary paths are forbidden")
        return parts

    @contextmanager
    def directory(self, parts, create=False):
        fd = os.dup(self.root)
        try:
            for part in parts:
                if create:
                    try:
                        os.mkdir(part, mode=0o755, dir_fd=fd)
                    except FileExistsError:
                        pass
                child = os.open(part, os.O_RDONLY | os.O_DIRECTORY | os.O_NOFOLLOW, dir_fd=fd)
                os.close(fd)
                fd = child
            yield fd
        finally:
            os.close(fd)

    @staticmethod
    def regular(fd, name, missing=False):
        try:
            info = os.stat(name, dir_fd=fd, follow_symlinks=False)
        except FileNotFoundError:
            if missing:
                return None
            raise
        if not stat.S_ISREG(info.st_mode) or info.st_nlink != 1:
            raise ValueError("Only ordinary files are allowed; symlinks and hardlinks are forbidden")
        return info

    def read(self, path):
        parts = self.parts(path)
        with self.directory(parts[:-1]) as parent:
            self.regular(parent, parts[-1])
            fd = os.open(parts[-1], os.O_RDONLY | os.O_NOFOLLOW | os.O_NONBLOCK, dir_fd=parent)
            with os.fdopen(fd, "rb") as f:
                info = os.fstat(f.fileno())
                if not stat.S_ISREG(info.st_mode) or info.st_nlink != 1:
                    raise ValueError("Only ordinary files are allowed")
                data = f.read(MAX_FILE + 1)
        if len(data) > MAX_FILE:
            raise ValueError("Text files are limited to 256 KiB")
        return data

    def save(self, path, data, overwrite=False, expected=None):
        if not isinstance(overwrite, bool):
            raise ValueError("overwrite must be true or false")
        parts = self.parts(path)
        with self.directory(parts[:-1], create=True) as parent:
            info = self.regular(parent, parts[-1], missing=True)
            if info and not overwrite:
                raise ValueError("File exists; read it first or explicitly set overwrite=true")
            if expected is not None and hashlib.sha256(self.read(path)).hexdigest() != expected:
                raise ValueError("File changed after reading; read it again before editing")
            temp = ".strata-tmp-" + uuid.uuid4().hex
            fd = os.open(temp, os.O_WRONLY | os.O_CREAT | os.O_EXCL | os.O_NOFOLLOW, 0o644, dir_fd=parent)
            try:
                with os.fdopen(fd, "wb") as f:
                    f.write(data)
                    f.flush()
                    os.fsync(f.fileno())
                if overwrite:
                    self.regular(parent, parts[-1], missing=True)
                    os.replace(temp, parts[-1], src_dir_fd=parent, dst_dir_fd=parent)
                else:
                    # Atomic create without clobbering a file created while generating.
                    os.link(temp, parts[-1], src_dir_fd=parent, dst_dir_fd=parent, follow_symlinks=False)
                return {"path": path, "bytes": len(data), "sha256": hashlib.sha256(data).hexdigest()}
            finally:
                try:
                    os.unlink(temp, dir_fd=parent)
                except FileNotFoundError:
                    pass

    def call(self, name, args):
        tool = next((t for t in TOOLS if t["name"] == name), None)
        if tool is None:
            raise ValueError("Unknown tool")
        if not isinstance(args, dict) or set(args) - set(tool["inputSchema"]["properties"]):
            raise ValueError("Arguments must be an object containing only documented fields")
        if any(k not in args for k in tool["inputSchema"]["required"]):
            raise ValueError("Missing required argument")
        path = args.get("path", ".")
        if name == "list_files":
            with self.directory(self.parts(path, directory=True)) as fd:
                entries = []
                with os.scandir(fd) as scan:
                    for entry in scan:
                        if len(entries) == 500:
                            return {"path": path, "entries": entries, "truncated": True}
                        if entry.name.startswith(".strata-tmp-"):
                            continue
                        info = entry.stat(follow_symlinks=False)
                        kind = "directory" if stat.S_ISDIR(info.st_mode) else "file" if stat.S_ISREG(info.st_mode) and info.st_nlink == 1 else "blocked"
                        entries.append({"name": entry.name, "type": kind, "bytes": info.st_size if kind == "file" else None})
            return {"path": path, "entries": sorted(entries, key=lambda x: x["name"]), "truncated": False}
        if name == "read_file":
            data = self.read(path)
            lines = data.decode("utf-8").splitlines()
            start, count = args.get("start_line", 1), args.get("max_lines", 100)
            if type(start) is not int or start < 1 or type(count) is not int or not 1 <= count <= 200:
                raise ValueError("start_line >= 1 and max_lines between 1 and 200 are required")
            selected, chars = [], 0
            for i in range(start - 1, min(len(lines), start - 1 + count)):
                line = f"{i + 1}: {lines[i]}"
                if len(line) > 12000:
                    raise ValueError("Line exceeds 12000 characters; use shorter source lines")
                if chars + len(line) > 14000:
                    break
                selected.append(line)
                chars += len(line) + 1
            next_line = start + len(selected)
            return {"path": path, "sha256": hashlib.sha256(data).hexdigest(), "total_lines": len(lines),
                    "content": "\n".join(selected), "next_line": next_line if next_line <= len(lines) else None}
        if name in ("write_file", "edit_file"):
            if name == "write_file":
                content = text_arg(args, "content")
                expected = None
                overwrite = args.get("overwrite", False)
            else:
                before = self.read(path)
                content = before.decode("utf-8")
                expected = args.get("expected_sha256", hashlib.sha256(before).hexdigest())
                if not isinstance(expected, str) or expected != hashlib.sha256(before).hexdigest():
                    raise ValueError("File changed after reading; read it again")
                old, new = text_arg(args, "old_text", nonempty=True), text_arg(args, "new_text")
                if content.count(old) != 1:
                    raise ValueError("old_text must match exactly once; include more surrounding code")
                content, overwrite = content.replace(old, new, 1), True
            data = content.encode("utf-8")
            if len(data) > MAX_FILE:
                raise ValueError("Text files are limited to 256 KiB")
            return self.save(path, data, overwrite, expected)
        return self.presentation(path, args)

    def presentation(self, path, args):
        self.parts(path)
        if not path.lower().endswith(".pptx"):
            raise ValueError("Output must end in .pptx (modern PowerPoint format)")
        title = text_arg(args, "title", 80, True)
        if any(c in title for c in "\n\r\t"):
            raise ValueError("Presentation title must be one paragraph without line breaks")
        slides = args["slides"]
        if not isinstance(slides, list) or not 1 <= len(slides) <= 40:
            raise ValueError("Supply between 1 and 40 slides")
        from pptx import Presentation
        from pptx.dml.color import RGBColor
        from pptx.util import Inches, Pt
        prs = Presentation()
        prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
        prs.core_properties.title = title
        prs.core_properties.author = "Strata"
        for number, spec in enumerate(slides, 1):
            if not isinstance(spec, dict) or set(spec) != {"title", "bullets"}:
                raise ValueError("Each slide requires exactly title and bullets")
            heading = text_arg(spec, "title", 40, True)
            bullets = spec["bullets"]
            if not isinstance(bullets, list) or len(bullets) > 6:
                raise ValueError("Each slide allows at most 6 bullets")
            for bullet in bullets:
                text_arg({"bullet": bullet}, "bullet", 120)
            if any(c in paragraph for paragraph in [heading, *bullets] for c in "\n\r\t"):
                raise ValueError("Slide title and each bullet must be one paragraph without line breaks")
            if len(heading) + sum(len(b) for b in bullets) > 360:
                raise ValueError("Slide exceeds 360 characters; split it into more slides")
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            slide.background.fill.solid()
            slide.background.fill.fore_color.rgb = RGBColor(248, 250, 252)
            top = slide.shapes.add_textbox(Inches(.7), Inches(.45), Inches(11.9), Inches(1.3)).text_frame
            top.word_wrap = True
            p = top.paragraphs[0]
            p.text = heading
            p.font.name, p.font.size, p.font.bold = "Microsoft JhengHei", Pt(32), True
            p.font.color.rgb = RGBColor(15, 23, 42)
            body = slide.shapes.add_textbox(Inches(.9), Inches(1.75), Inches(11.5), Inches(4.95)).text_frame
            body.word_wrap = True
            for i, bullet in enumerate(bullets):
                p = body.paragraphs[0] if i == 0 else body.add_paragraph()
                p.text = "• " + bullet
                p.font.name, p.font.size = "Microsoft JhengHei", Pt(20)
                p.font.color.rgb = RGBColor(51, 65, 85)
                p.space_after, p.line_spacing = Pt(12), 1.15
            foot = slide.shapes.add_textbox(Inches(.7), Inches(7), Inches(11.9), Inches(.3)).text_frame.paragraphs[0]
            foot.text = f"{title[:60]}  {number}/{len(slides)}"
            foot.font.name, foot.font.size = "Microsoft JhengHei", Pt(10)
            foot.font.color.rgb = RGBColor(100, 116, 139)
        output = io.BytesIO()
        prs.save(output)
        data = output.getvalue()
        if len(data) > MAX_PPTX:
            raise ValueError("Presentation exceeds 8 MiB")
        return {**self.save(path, data, args.get("overwrite", False)), "slides": len(slides), "format": "pptx"}


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--root", required=True)
    options = parser.parse_args()
    workspace = Workspace(options.root)
    try:
        while True:
            raw = sys.stdin.buffer.readline(MAX_MESSAGE + 1)
            if not raw:
                break
            request_id = None
            try:
                if len(raw) > MAX_MESSAGE:
                    while raw and not raw.endswith(b"\n"):
                        raw = sys.stdin.buffer.readline(MAX_MESSAGE + 1)
                    raise ValueError("Message exceeds 2 MiB")
                req = json.loads(raw)
                if not isinstance(req, dict):
                    raise ValueError("JSON-RPC request must be an object")
                request_id = req.get("id")
                if req.get("jsonrpc") != "2.0" or not isinstance(req.get("method"), str):
                    raise ValueError("Invalid JSON-RPC request")
                if "id" not in req:
                    continue
                method, params = req["method"], req.get("params", {})
                if not isinstance(params, dict):
                    raise ValueError("params must be an object")
                if method == "initialize":
                    result = {"protocolVersion": "2025-06-18", "capabilities": {"tools": {}},
                              "serverInfo": {"name": "strata-workspace", "version": "1.0"},
                              "instructions": "Tools access only /workspace. Use relative paths. Read before edits. Save actual files when requested; report returned path."}
                elif method == "ping":
                    result = {}
                elif method == "tools/list":
                    result = {"tools": TOOLS}
                elif method == "tools/call":
                    try:
                        data = workspace.call(params.get("name"), params.get("arguments", {}))
                        result = {"content": [{"type": "text", "text": json.dumps(data, ensure_ascii=False)}]}
                    except (ValueError, OSError, KeyError, TypeError, UnicodeError) as exc:
                        result = {"isError": True, "content": [{"type": "text", "text": str(exc)}]}
                else:
                    print(json.dumps({"jsonrpc": "2.0", "id": request_id, "error": {"code": -32601, "message": "Method not found"}}), flush=True)
                    continue
                response = {"jsonrpc": "2.0", "id": request_id, "result": result}
            except (ValueError, TypeError) as exc:
                response = {"jsonrpc": "2.0", "id": request_id, "error": {"code": -32600, "message": str(exc)}}
            print(json.dumps(response, ensure_ascii=False), flush=True)
    finally:
        workspace.close()


if __name__ == "__main__":
    main()
